import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

doc_path = "Plantila para proyectos EIS Albarregas.docx"
img_path = "Mapa_de_empatia_FacturaMetrica.png"

doc = docx.Document(doc_path)

# Find the insertion point: Right after "Fase de Análisis" (Heading 1) and before "Requisitos Funcionales de la App"
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Requisitos Funcionales de la App":
        target_idx = i
        break

if target_idx is None:
    print("Could not find 'Requisitos Funcionales de la App'!")
    exit(1)

print(f"Target index found at paragraph {target_idx}: '{doc.paragraphs[target_idx].text}'")

ref_p = doc.paragraphs[target_idx]

def insert_p_before(ref, text, style=None, bold_prefix=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    new_p = ref.insert_paragraph_before()
    if style:
        new_p.style = doc.styles[style] if style in doc.styles else doc.styles['Normal']
    new_p.alignment = align
    if bold_prefix:
        r_bold = new_p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.name = "Calibri"
    if text:
        r_text = new_p.add_run(text)
        r_text.font.name = "Calibri"
    return new_p

# Add section title
h2 = insert_p_before(ref_p, "Técnicas de Análisis y Empatía de Usuario (Design Thinking: Mapa de Empatía)", style="Heading 2")

insert_p_before(ref_p, 
    "Con el objetivo de fundamentar los requisitos del sistema en problemas y necesidades humanas reales, se ha aplicado la metodología de Design Thinking mediante la construcción del Mapa de Empatía (modelo formal XPLANE). FacturaMétrica opera en un modelo B2B2C donde convergen dos perfiles interdependientes: el profesional autónomo (usuario directo de la app móvil) y el gestor/contable (cliente corporativo del panel web).",
    style="Normal"
)

# Insert the diagram image
p_img = ref_p.insert_paragraph_before()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_img = p_img.add_run()
run_img.add_picture(img_path, width=Inches(6.2))

# Caption
p_caption = ref_p.insert_paragraph_before()
p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_cap = p_caption.add_run("Figura 1: Mapa de Empatía del usuario arquetipo (Autónomo / Contribuyente) para FacturaMétrica.")
r_cap.font.size = Pt(9.5)
r_cap.font.italic = True
r_cap.font.color.rgb = RGBColor(100, 116, 139)

insert_p_before(ref_p,
    "El desglose detallado de las seis dimensiones del mapa de empatía revela las siguientes conclusiones clave:",
    style="Normal"
)

# Bullet points for the 6 sections
insert_p_before(ref_p, 
    " El autónomo vive enfocado en generar negocio y atender clientes. La recopilación manual de comprobantes le produce una profunda aversión burocrática. Su mayor inquietud es el miedo a sanciones de la Agencia Tributaria (AEAT) por deducir gastos en fines de semana o suministros sin saber si cumplen los criterios normativos de la LIRPF. Aspira a la tranquilidad total y al control sin fricciones.",
    style="Normal",
    bold_prefix="• 1. ¿Qué Piensa y Siente? (Inquietudes y aspiraciones):"
)

insert_p_before(ref_p, 
    " Escucha continuamente a su asesor recordándole la obligación de aportar facturas completas antes del día 15 y advirtiéndole de que los tickets simplificados o no justificados acarrearán inspecciones. En su entorno profesional escucha anécdotas de inspecciones paralelas y noticias sobre la implantación de la Ley Antifraude y el sistema Veri*Factu.",
    style="Normal",
    bold_prefix="• 2. ¿Qué Oye? (Influencias y entorno):"
)

insert_p_before(ref_p, 
    " En su día a día observa tickets de papel térmico acumulados en la guantera del vehículo o carteras, cuya tinta se borra antes del cierre de trimestre. En el mercado observa herramientas contables excesivamente densas o aplicaciones de escaneo que no aportan ningún valor fiscal preventivo.",
    style="Normal",
    bold_prefix="• 3. ¿Qué Ve? (Realidad física y oferta de mercado):"
)

insert_p_before(ref_p, 
    " En restaurantes y comercios pide el ticket con prisa con la promesa de revisarlo más tarde, olvidándolo con frecuencia. En vísperas del cierre trimestral (enero, abril, julio, octubre) sufre fines de semana maratonianos buscando comprobantes y enviando fotografías torcidas por WhatsApp a su despacho contable.",
    style="Normal",
    bold_prefix="• 4. ¿Qué Dice y Hace? (Conducta y hábitos cotidianos):"
)

insert_p_before(ref_p, 
    " Pánico a requerimientos tributarios, liquidaciones paralelas y recargos por deducciones denegadas; pérdida directa de dinero al no deducir gastos legítimos extraviados o descoloridos; y tensiones recurrentes con su gestor debido a envíos tardíos e incompletos.",
    style="Normal",
    bold_prefix="• 5. Esfuerzos (Frustraciones, dolores y barreras):"
)

insert_p_before(ref_p, 
    " Captura inmediata en 3 segundos desde el smartphone con OCR instantáneo; asistencia fiscal interactiva ('Asesor Fiscal Virtual') que le avisa en el instante preciso para justificar gastos de riesgo; y paz mental derivada de saber que sus comprobantes están validados y volcados a la contabilidad del despacho.",
    style="Normal",
    bold_prefix="• 6. Resultados (Deseos de éxito y beneficios alcanzados):"
)

insert_p_before(ref_p,
    "Matriz de Trazabilidad: De los Puntos de Dolor (Esfuerzos) a los Requisitos del Sistema",
    style="Normal",
    bold_prefix="Tabla de Alineamiento: "
)

# Create a clean Word table mapping Pain Points to Software Requirements
# We insert table before ref_p
new_table = doc.add_table(rows=5, cols=3)
new_table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["Dimensión / Dolor Detectado (Design Thinking)", "Respuesta de FacturaMétrica", "Requisito Asociado"]
hdr_cells = new_table.rows[0].cells
for j, text in enumerate(headers):
    hdr_cells[j].text = text
    # Set background shading and bold
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1E293B')
    hdr_cells[j]._tc.get_or_add_tcPr().append(shd)
    for p in hdr_cells[j].paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

table_data = [
    (
        "Miedo a sanciones e inspecciones de la AEAT por deducir gastos de fin de semana o comidas no acreditadas.",
        "Asesor Fiscal Virtual con semáforo preventivo que detecta horarios atípicos y solicita justificación contextual inmediata.",
        "RF-03, RF-04\n(Validación Contextual)"
    ),
    (
        "Pérdida de comprobantes en papel térmico borrados por calor o extraviados en carteras.",
        "Captura fotográfica inmediata en 3 segundos con almacenamiento seguro en la nube y repositorio documental indexado.",
        "RF-01, RF-12\n(Captura & Repositorio)"
    ),
    (
        "Pérdida de horas de trabajo en cierres trimestrales y envío caótico de fotos por WhatsApp.",
        "Sincronización automática con el panel del gestor, clasificación semafórica y aprobación de criterios en 1 solo clic.",
        "RF-05, RF-06\n(Bandeja Gestor)"
    ),
    (
        "Sobrecarga del contable picando datos a mano en su programa de contabilidad (A3, Sage, Wolters).",
        "Asignación automática de subcuentas del PGC (629) y exportación directa de ficheros estructurados.",
        "RF-07, RF-08\n(Integración ERP)"
    )
]

for row_idx, data in enumerate(table_data, start=1):
    row_cells = new_table.rows[row_idx].cells
    for col_idx, cell_value in enumerate(data):
        row_cells[col_idx].text = cell_value
        for p in row_cells[col_idx].paragraphs:
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9)
        # zebra striping
        if row_idx % 2 == 1:
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F1F5F9')
            row_cells[col_idx]._tc.get_or_add_tcPr().append(shd)

# Move the newly added table to immediately before ref_p in the XML tree
table_element = new_table._tbl
ref_element = ref_p._p
ref_element.addprevious(table_element)

# Add spacing paragraph before Requisitos Funcionales
spacer_p = ref_p.insert_paragraph_before()
spacer_p.text = ""

doc.save("Plantila para proyectos EIS Albarregas.docx")
print("Successfully integrated Empathy Map into 'Plantila para proyectos EIS Albarregas.docx'!")
